"""Build the developer/BA Word handover from reviewed Markdown; resolve source links."""
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT=Path(__file__).resolve().parents[1]
sources={p.stem:p for p in (ROOT/'src/main/java').rglob('*.java')}
def source_ref(match):
    cls,method=match.group(1).split('.')
    path=sources[cls]
    lines=path.read_text(encoding='utf-8').splitlines()
    candidates=[i+1 for i,line in enumerate(lines) if re.search(r'^\s*(?:(?:public|private|protected|static|final|synchronized)\s+)*[\w<>?,.\[\] ]+\s+'+re.escape(method)+r'\s*\(',line) and not line.lstrip().startswith(('return ','if ','var '))]
    if not candidates: raise ValueError('Unresolved method '+match.group(1))
    line=candidates[0]
    return f'[{cls}.{method} dòng {line}]({path.as_posix()}:{line})'
for filename in ['LLM_FLOW_DEV_BA.md','PRODUCTION_DATA_FLOW_REVIEW.md']:
    path=ROOT/'docs'/filename
    text=path.read_text(encoding='utf-8')
    text=re.sub(r'\{@([\w]+\.[\w]+)\}',source_ref,text)
    path.write_text(text,encoding='utf-8')

document=Document()
# The bundled default template can carry a blue Title border; remove paragraph borders.
for border in list(document.styles.element.iter(qn('w:pBdr'))):
    border.getparent().remove(border)
section=document.sections[0]
section.page_width=Inches(8.5); section.page_height=Inches(11)
section.top_margin=Inches(.72); section.bottom_margin=Inches(.7)
section.left_margin=Inches(.72); section.right_margin=Inches(.72)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','List Bullet']:
    style=document.styles[name]
    style.font.name='Arial'; style.font.color.rgb=RGBColor(0,0,0)
    style.font.size=Pt(11)
    style.paragraph_format.space_after=Pt(6)
    style.paragraph_format.line_spacing=1.12
document.styles['Title'].font.size=Pt(25)
document.styles['Title'].paragraph_format.space_after=Pt(14)
for name,size in [('Heading 1',16),('Heading 2',12.5),('Heading 3',11.5)]:
    document.styles[name].font.size=Pt(size)
    document.styles[name].font.bold=True
    document.styles[name].paragraph_format.space_before=Pt(14)
    document.styles[name].paragraph_format.keep_with_next=True

def inline(paragraph,text):
    pattern=r'(\[[^\]]+\]\([^\)]+\)|\*\*[^*]+\*\*|`[^`]+`)'
    for part in re.split(pattern,text):
        link=re.fullmatch(r'\[([^\]]+)\]\(([^\)]+)\)',part)
        if link:
            label,url=link.groups()
            if re.match(r'^[A-Za-z]:/',url):
                url='file:///'+url.rsplit(':',1)[0]
            hyperlink=OxmlElement('w:hyperlink')
            hyperlink.set(qn('r:id'),paragraph.part.relate_to(url,RT.HYPERLINK,is_external=True))
            run=OxmlElement('w:r'); pr=OxmlElement('w:rPr'); color=OxmlElement('w:color');color.set(qn('w:val'),'174A70');pr.append(color);run.append(pr)
            value=OxmlElement('w:t');value.text=label;run.append(value);hyperlink.append(run);paragraph._p.append(hyperlink)
        else:
            run=paragraph.add_run(part[2:-2] if part.startswith('**') else part[1:-1] if part.startswith('`') else part)
            if part.startswith('**'):run.bold=True

def table(rows):
    cols=len(rows[0]); tbl=document.add_table(rows=0, cols=cols)
    tbl.alignment=WD_TABLE_ALIGNMENT.CENTER; tbl.autofit=False
    widths=([2.55,4.51] if cols==2 and rows[0][0]=='Luật' else [2.12,4.94] if cols==2 else [1.95,2.2,2.91] if cols==3 else [1.7,1.45,1.97,1.94])
    if rows[0][0]=='Nhóm API': widths=[.82,3.65,2.59]
    for col,width in zip(tbl.columns,widths):col.width=Inches(width)
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
    tbl._tbl.tblPr.append(borders)
    for ri,data in enumerate(rows):
        row=tbl.add_row()
        if ri==0:
            header=OxmlElement('w:tblHeader');row._tr.get_or_add_trPr().append(header)
        # Keep one logical record together; all rows are shorter than a page.
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for ci,(cell,value) in enumerate(zip(row.cells,data)):
            cell.width=Inches(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=cell._tc.get_or_add_tcPr(); margins=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                el=OxmlElement('w:'+edge);el.set(qn('w:w'),'90');el.set(qn('w:type'),'dxa');margins.append(el)
            tcpr.append(margins)
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'173B52' if ri==0 else 'F1F5F8' if ri%2==0 else 'FFFFFF');tcpr.append(shade)
            p=cell.paragraphs[0];p.paragraph_format.line_spacing=1.05;p.paragraph_format.space_after=Pt(2)
            if ri==0:p.paragraph_format.keep_with_next=True
            inline(p,value)
            for r in p.runs:
                r.font.size=Pt(10)
                if ri==0:r.bold=True;r.font.color.rgb=RGBColor(255,255,255)
    document.add_paragraph().paragraph_format.space_after=Pt(2)

lines=(ROOT/'docs/LLM_FLOW_DEV_BA.md').read_text(encoding='utf-8').splitlines()
i=0
while i<len(lines):
    line=lines[i].strip()
    if not line:i+=1;continue
    if line.startswith('|'):
        rows=[]
        while i<len(lines) and lines[i].strip().startswith('|'):
            row=[s.strip() for s in lines[i].strip().strip('|').split('|')]
            if not all(re.fullmatch(r'[:\- ]+',s) for s in row):rows.append(row)
            i+=1
        table(rows);continue
    if line.startswith('# '):
        document.add_paragraph(line[2:],style='Title')
    elif line.startswith('## '):
        document.add_paragraph(line[3:],style='Heading 1')
    elif line.startswith('### '):
        document.add_paragraph(line[4:],style='Heading 2')
    elif line.startswith('- '):
        inline(document.add_paragraph(style='List Bullet'),line[2:])
    else:inline(document.add_paragraph(),line)
    i+=1
footer=section.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('Luồng LLM  |  Trang ').font.size=Pt(9)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
document.core_properties.title='Luồng xử lý LLM cho tin tức và tài chính'
document.core_properties.subject='Bàn giao luồng code và hợp đồng dữ liệu cho Dev và BA'
document.core_properties.author='Dự án Hệ thống dữ liệu tài chính'
output=ROOT/'docs/BAO_CAO_LUONG_LLM_DEV_BA.docx'
document.save(output)
print(output)
print('paragraphs',len(document.paragraphs),'tables',len(document.tables))
